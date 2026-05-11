"""
SkillTrack System — Desktop Admin Dashboard
Python 3 + Tkinter + mysql-connector-python

Run:
    pip install mysql-connector-python openpyxl
    python app.py
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import mysql.connector
from mysql.connector import Error
import csv
import os

try:
    import openpyxl
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

# ─── Theme / Colour Palette ──────────────────────────────────────────────────
BG_DARK      = "#0f1117"
BG_PANEL     = "#161b27"
BG_SIDEBAR   = "#0d1120"
ACCENT       = "#4f8ef7"
TEXT_MAIN    = "#e8eaf0"
TEXT_MUTED   = "#6b7280"
TEXT_HEADER  = "#ffffff"
ROW_ODD      = "#161b27"
ROW_EVEN     = "#1a2035"
ROW_SEL      = "#1e3a6e"
BTN_ADD      = "#22c55e"
BTN_UPD      = "#f59e0b"
BTN_DEL      = "#ef4444"
BTN_REF      = "#4f8ef7"
BTN_EXP      = "#8b5cf6"
BORDER       = "#252d42"
CARD_BG      = "#1a2035"

FONT_TITLE   = ("Segoe UI", 18, "bold")
FONT_SUB     = ("Segoe UI", 11, "bold")
FONT_BODY    = ("Segoe UI", 10)
FONT_SMALL   = ("Segoe UI", 9)
FONT_MONO    = ("Consolas", 10)
FONT_CARD_N  = ("Segoe UI", 28, "bold")
FONT_CARD_L  = ("Segoe UI", 10)

TABLES = [
    "learners",
    "instructors",
    "courses",
    "enrollments",
    "progress",
    "certificates",
    "attendance",
]

DB_CONFIG = dict(host="localhost", user="root", password="", database="skilltrack")


# ─── Helper ──────────────────────────────────────────────────────────────────

def styled_button(parent, text, color, command, width=12):
    btn = tk.Button(
        parent, text=text, command=command,
        bg=color, fg="#ffffff", activebackground=color,
        activeforeground="#ffffff", relief="flat", bd=0,
        font=FONT_BODY, cursor="hand2", width=width,
        padx=8, pady=6,
    )
    return btn


# ─── Add / Edit Dialog ───────────────────────────────────────────────────────

class RecordDialog(tk.Toplevel):
    def __init__(self, parent, title, columns, values=None, pk_col=None):
        super().__init__(parent)
        self.title(title)
        self.configure(bg=BG_DARK)
        self.resizable(False, False)
        self.grab_set()

        self.result = None
        self.entries = {}

        tk.Label(self, text=title, font=FONT_SUB, bg=BG_DARK,
                 fg=TEXT_HEADER).pack(pady=(18, 10), padx=24, anchor="w")

        frame = tk.Frame(self, bg=BG_DARK)
        frame.pack(padx=24, pady=4)

        for i, col in enumerate(columns):
            is_pk = (col == pk_col)
            tk.Label(frame, text=col, font=FONT_SMALL, bg=BG_DARK,
                     fg=TEXT_MUTED).grid(row=i, column=0, sticky="w",
                                         pady=4, padx=(0, 12))
            var = tk.StringVar(value=(values[i] if values else ""))
            state = "disabled" if is_pk else "normal"
            ent = tk.Entry(frame, textvariable=var, font=FONT_MONO,
                           bg=BG_PANEL, fg=TEXT_MAIN,
                           insertbackground=TEXT_MAIN,
                           disabledbackground=BG_SIDEBAR,
                           disabledforeground=TEXT_MUTED,
                           relief="flat", bd=0, highlightthickness=1,
                           highlightbackground=BORDER,
                           highlightcolor=ACCENT,
                           width=34, state=state)
            ent.grid(row=i, column=1, pady=4, ipady=5)
            self.entries[col] = var

        btn_frame = tk.Frame(self, bg=BG_DARK)
        btn_frame.pack(pady=(16, 20), padx=24, anchor="e")
        styled_button(btn_frame, "Cancel", TEXT_MUTED,
                      self.destroy, width=10).pack(side="right", padx=(6, 0))
        styled_button(btn_frame, "Confirm", ACCENT,
                      self._confirm, width=10).pack(side="right")

        self.center(parent)

    def center(self, parent):
        self.update_idletasks()
        pw, ph = parent.winfo_width(), parent.winfo_height()
        px, py = parent.winfo_x(), parent.winfo_y()
        w, h = self.winfo_width(), self.winfo_height()
        self.geometry(f"{w}x{h}+{px + pw//2 - w//2}+{py + ph//2 - h//2}")

    def _confirm(self):
        self.result = {col: var.get() for col, var in self.entries.items()}
        self.destroy()


# ─── Dashboard Window ────────────────────────────────────────────────────────

class DashboardWindow(tk.Toplevel):
    def __init__(self, parent, cursor):
        super().__init__(parent)
        self.title("SkillTrack - Dashboard")
        self.configure(bg=BG_DARK)
        self.geometry("780x520")
        self.resizable(False, False)
        self.grab_set()
        self.cursor = cursor
        self._build(parent)

    def _build(self, parent):
        tk.Label(self, text="Dashboard", font=FONT_TITLE,
                 bg=BG_DARK, fg=TEXT_HEADER).pack(anchor="w", padx=28, pady=(22, 4))
        tk.Label(self, text="Live stats from your database",
                 font=FONT_SMALL, bg=BG_DARK, fg=TEXT_MUTED).pack(anchor="w", padx=28)

        cards_frame = tk.Frame(self, bg=BG_DARK)
        cards_frame.pack(fill="x", padx=20, pady=18)

        stats = self._get_stats()
        labels = ["Learners", "Instructors", "Courses", "Enrollments",
                  "Certificates", "Avg Progress %", "Attendance Records"]
        colors = [ACCENT, BTN_ADD, BTN_UPD, BTN_DEL, BTN_EXP, "#06b6d4", "#ec4899"]

        for i, (label, value, color) in enumerate(zip(labels, stats, colors)):
            card = tk.Frame(cards_frame, bg=CARD_BG, width=158, height=90)
            card.grid(row=i // 4, column=i % 4, padx=6, pady=6)
            card.grid_propagate(False)
            tk.Label(card, text=str(value), font=FONT_CARD_N,
                     bg=CARD_BG, fg=color).place(relx=0.5, rely=0.38, anchor="center")
            tk.Label(card, text=label, font=FONT_CARD_L,
                     bg=CARD_BG, fg=TEXT_MUTED).place(relx=0.5, rely=0.75, anchor="center")

        tk.Label(self, text="Top Learners by Progress",
                 font=FONT_SUB, bg=BG_DARK, fg=TEXT_HEADER).pack(anchor="w", padx=28, pady=(8, 4))

        frame = tk.Frame(self, bg=BG_DARK)
        frame.pack(fill="both", expand=True, padx=20, pady=(0, 16))

        style = ttk.Style()
        style.configure("Dash.Treeview", background=ROW_ODD, foreground=TEXT_MAIN,
                        fieldbackground=ROW_ODD, rowheight=26, font=FONT_BODY)
        style.configure("Dash.Treeview.Heading", background=BG_SIDEBAR,
                        foreground=ACCENT, font=("Segoe UI", 9, "bold"))

        tree = ttk.Treeview(frame, style="Dash.Treeview", show="headings",
                            columns=("name", "course", "progress"))
        tree.heading("name",     text="Learner")
        tree.heading("course",   text="Course")
        tree.heading("progress", text="Progress %")
        tree.column("name",     width=220)
        tree.column("course",   width=220)
        tree.column("progress", width=120, anchor="center")
        tree.pack(fill="both", expand=True)

        rows = self._get_top_learners()
        for i, row in enumerate(rows):
            tag = "even" if i % 2 == 0 else "odd"
            tree.insert("", "end", values=row, tags=(tag,))
        tree.tag_configure("odd",  background=ROW_ODD)
        tree.tag_configure("even", background=ROW_EVEN)

        styled_button(self, "Close", TEXT_MUTED, self.destroy, 10).pack(pady=(0, 14))

    def _get_stats(self):
        results = []
        queries = [
            "SELECT COUNT(*) FROM learners",
            "SELECT COUNT(*) FROM instructors",
            "SELECT COUNT(*) FROM courses",
            "SELECT COUNT(*) FROM enrollments",
            "SELECT COUNT(*) FROM certificates",
            "SELECT ROUND(AVG(completion_pct),1) FROM progress",
            "SELECT COUNT(*) FROM attendance",
        ]
        for q in queries:
            try:
                self.cursor.execute(q)
                val = self.cursor.fetchone()[0]
                results.append(val if val is not None else 0)
            except Exception:
                results.append("-")
        return results

    def _get_top_learners(self):
        try:
            self.cursor.execute("""
                SELECT l.full_name, c.title, p.completion_pct
                FROM progress p
                JOIN enrollments e ON p.enrollment_id = e.enrollment_id
                JOIN learners l    ON e.learner_id    = l.learner_id
                JOIN courses  c    ON e.course_id     = c.course_id
                ORDER BY p.completion_pct DESC
                LIMIT 10
            """)
            return self.cursor.fetchall()
        except Exception:
            return []


# ─── Main Application ─────────────────────────────────────────────────────────

class SkillTrackApp(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("SkillTrack - Admin Dashboard")
        self.geometry("1280x740")
        self.minsize(900, 580)
        self.configure(bg=BG_DARK)

        self.conn            = None
        self.cursor          = None
        self.current_table   = tk.StringVar(value=TABLES[0])
        self.columns         = []
        self.all_rows        = []
        self.status_text     = tk.StringVar(value="Connecting ...")
        self._sidebar_btns   = {}
        self._search_var     = tk.StringVar()
        self._search_var.trace("w", lambda *a: self._apply_search())

        self._build_ui()
        self.connect_to_db()
        self.load_table(TABLES[0])

    # ── DB ────────────────────────────────────────────────────────────────────

    def connect_to_db(self):
        try:
            if self.conn and self.conn.is_connected():
                self.conn.close()
            self.conn   = mysql.connector.connect(**DB_CONFIG)
            self.cursor = self.conn.cursor()
            self.set_status("Connected to MySQL - skilltrack", ok=True)
        except Error as e:
            self.set_status(f"Connection failed: {e}", ok=False)
            messagebox.showerror("DB Error",
                f"Cannot connect to MySQL.\n\n{e}\n\n"
                "Check that MySQL is running and credentials are correct in app.py.")

    def ensure_connection(self):
        try:
            if self.conn is None or not self.conn.is_connected():
                self.connect_to_db()
        except Exception:
            self.connect_to_db()

    # ── UI Build ──────────────────────────────────────────────────────────────

    def _build_ui(self):
        # Sidebar
        sidebar = tk.Frame(self, bg=BG_SIDEBAR, width=200)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)

        logo_frame = tk.Frame(sidebar, bg=BG_SIDEBAR)
        logo_frame.pack(fill="x", pady=(20, 6), padx=16)
        tk.Label(logo_frame, text="ST", font=("Segoe UI", 16, "bold"),
                 bg=ACCENT, fg="white", width=3).pack(side="left")
        tk.Label(logo_frame, text=" SkillTrack", font=("Segoe UI", 13, "bold"),
                 bg=BG_SIDEBAR, fg=TEXT_HEADER).pack(side="left")

        tk.Frame(sidebar, bg=BORDER, height=1).pack(fill="x", padx=16, pady=(12, 6))
        styled_button(sidebar, "Dashboard", BTN_EXP,
                      self.open_dashboard, width=18).pack(padx=12, pady=4)
        tk.Frame(sidebar, bg=BORDER, height=1).pack(fill="x", padx=16, pady=(6, 10))

        tk.Label(sidebar, text="TABLES", font=("Segoe UI", 8, "bold"),
                 bg=BG_SIDEBAR, fg=TEXT_MUTED).pack(anchor="w", padx=20, pady=(4, 6))

        for tbl in TABLES:
            btn = tk.Button(
                sidebar, text=f"  {tbl}", anchor="w",
                font=FONT_BODY, relief="flat", bd=0,
                bg=BG_SIDEBAR, fg=TEXT_MAIN, activebackground=BG_PANEL,
                activeforeground=ACCENT, cursor="hand2", pady=8,
                command=lambda t=tbl: self.load_table(t)
            )
            btn.pack(fill="x", padx=8)
            self._sidebar_btns[tbl] = btn

        tk.Frame(sidebar, bg=BORDER, height=1).pack(fill="x", padx=16, pady=(20, 10))
        styled_button(sidebar, "Reconnect DB", TEXT_MUTED,
                      self.connect_to_db, width=18).pack(padx=12, pady=4)

        # Main area
        main = tk.Frame(self, bg=BG_DARK)
        main.pack(side="left", fill="both", expand=True)

        # Top bar
        topbar = tk.Frame(main, bg=BG_PANEL, height=58)
        topbar.pack(fill="x")
        topbar.pack_propagate(False)

        self.lbl_table = tk.Label(topbar, text="", font=FONT_TITLE,
                                  bg=BG_PANEL, fg=TEXT_HEADER)
        self.lbl_table.pack(side="left", padx=24, pady=10)

        btn_frame = tk.Frame(topbar, bg=BG_PANEL)
        btn_frame.pack(side="right", padx=16, pady=10)

        styled_button(btn_frame, "Export",   BTN_EXP, self.export_data,   10).pack(side="left", padx=3)
        styled_button(btn_frame, "+ Add",    BTN_ADD, self.add_record,    10).pack(side="left", padx=3)
        styled_button(btn_frame, "Update",   BTN_UPD, self.update_record, 10).pack(side="left", padx=3)
        styled_button(btn_frame, "Delete",   BTN_DEL, self.delete_record, 10).pack(side="left", padx=3)
        styled_button(btn_frame, "Refresh",  BTN_REF, self.refresh,       10).pack(side="left", padx=3)

        # Search bar
        search_frame = tk.Frame(main, bg=BG_DARK)
        search_frame.pack(fill="x", padx=18, pady=(10, 0))

        tk.Label(search_frame, text="Search:", font=FONT_BODY,
                 bg=BG_DARK, fg=TEXT_MUTED).pack(side="left", padx=(0, 6))
        search_entry = tk.Entry(search_frame, textvariable=self._search_var,
                                font=FONT_BODY, bg=BG_PANEL, fg=TEXT_MAIN,
                                insertbackground=TEXT_MAIN, relief="flat", bd=0,
                                highlightthickness=1, highlightbackground=BORDER,
                                highlightcolor=ACCENT, width=40)
        search_entry.pack(side="left", ipady=6, padx=(0, 8))

        self.lbl_count = tk.Label(search_frame, text="", font=FONT_SMALL,
                                  bg=BG_DARK, fg=TEXT_MUTED)
        self.lbl_count.pack(side="left", padx=8)

        # Data grid
        grid_frame = tk.Frame(main, bg=BG_DARK)
        grid_frame.pack(fill="both", expand=True, padx=18, pady=10)

        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("SkillTrack.Treeview",
                        background=ROW_ODD, foreground=TEXT_MAIN,
                        fieldbackground=ROW_ODD, rowheight=28,
                        font=FONT_BODY, borderwidth=0)
        style.configure("SkillTrack.Treeview.Heading",
                        background=BG_SIDEBAR, foreground=ACCENT,
                        font=("Segoe UI", 9, "bold"), relief="flat", borderwidth=0)
        style.map("SkillTrack.Treeview",
                  background=[("selected", ROW_SEL)],
                  foreground=[("selected", "#ffffff")])

        self.tree = ttk.Treeview(grid_frame, style="SkillTrack.Treeview",
                                  selectmode="browse", show="headings")
        self.tree.tag_configure("odd",  background=ROW_ODD)
        self.tree.tag_configure("even", background=ROW_EVEN)

        vsb = ttk.Scrollbar(grid_frame, orient="vertical",   command=self.tree.yview)
        hsb = ttk.Scrollbar(grid_frame, orient="horizontal",  command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        self.tree.grid(row=0, column=0, sticky="nsew")
        grid_frame.rowconfigure(0, weight=1)
        grid_frame.columnconfigure(0, weight=1)

        # Status bar
        statusbar = tk.Frame(main, bg=BG_SIDEBAR, height=28)
        statusbar.pack(fill="x", side="bottom")
        statusbar.pack_propagate(False)
        self.lbl_status = tk.Label(statusbar, textvariable=self.status_text,
                                   font=FONT_SMALL, bg=BG_SIDEBAR, fg=TEXT_MUTED)
        self.lbl_status.pack(side="left", padx=14, pady=4)

    # ── Search ────────────────────────────────────────────────────────────────

    def _apply_search(self):
        query = self._search_var.get().strip().lower()

        for row in self.tree.get_children():
            self.tree.delete(row)

        filtered = [r for r in self.all_rows
                    if query == "" or any(query in str(v).lower() for v in r)]

        for i, row in enumerate(filtered):
            tag = "even" if i % 2 == 0 else "odd"
            self.tree.insert("", "end", values=row, tags=(tag,))

        total = len(self.all_rows)
        shown = len(filtered)
        self.lbl_count.config(
            text=f"{shown} of {total} rows" if query else f"{total} rows")

    # ── Table loading ─────────────────────────────────────────────────────────

    def load_table(self, table):
        self.ensure_connection()
        if not self.conn:
            return
        self.current_table.set(table)
        self.lbl_table.config(text=table)
        self._search_var.set("")

        for t, btn in self._sidebar_btns.items():
            btn.config(bg=BG_PANEL if t == table else BG_SIDEBAR,
                       fg=ACCENT   if t == table else TEXT_MAIN)
        try:
            self.cursor.execute(f"DESCRIBE `{table}`")
            self.columns = [row[0] for row in self.cursor.fetchall()]

            self.tree["columns"] = self.columns
            for col in self.columns:
                self.tree.heading(col, text=col)
                self.tree.column(col, width=130, anchor="w", minwidth=80)

            self._fetch_rows()
            self.set_status(f"Table: {table}  |  {len(self.all_rows)} rows", ok=True)
        except Error as e:
            messagebox.showerror("Query Error", str(e))

    def _fetch_rows(self):
        for row in self.tree.get_children():
            self.tree.delete(row)
        table = self.current_table.get()
        self.cursor.execute(f"SELECT * FROM `{table}`")
        self.all_rows = self.cursor.fetchall()
        for i, row in enumerate(self.all_rows):
            tag = "even" if i % 2 == 0 else "odd"
            self.tree.insert("", "end", values=row, tags=(tag,))
        self.lbl_count.config(text=f"{len(self.all_rows)} rows")

    def refresh(self):
        self.load_table(self.current_table.get())

    # ── CRUD ──────────────────────────────────────────────────────────────────

    def add_record(self):
        visible = self.columns[1:]
        dlg = RecordDialog(self, f"Add - {self.current_table.get()}", visible)
        self.wait_window(dlg)
        if dlg.result is None:
            return
        placeholders = ", ".join(["%s"] * len(visible))
        cols_str     = ", ".join([f"`{c}`" for c in visible])
        sql = f"INSERT INTO `{self.current_table.get()}` ({cols_str}) VALUES ({placeholders})"
        values = [dlg.result[c] or None for c in visible]
        try:
            self.cursor.execute(sql, values)
            self.conn.commit()
            self.refresh()
            self.set_status("Record added.", ok=True)
        except Error as e:
            messagebox.showerror("Insert Error", str(e))

    def _selected_values(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Select a row", "Please select a row first.")
            return None
        return self.tree.item(sel[0], "values")

    def update_record(self):
        values = self._selected_values()
        if values is None:
            return
        pk_col = self.columns[0]
        dlg = RecordDialog(self, f"Update - {self.current_table.get()}",
                           self.columns, list(values), pk_col=pk_col)
        self.wait_window(dlg)
        if dlg.result is None:
            return
        pk_val  = values[0]
        updates = ", ".join([f"`{c}` = %s" for c in self.columns[1:]])
        sql = f"UPDATE `{self.current_table.get()}` SET {updates} WHERE `{pk_col}` = %s"
        new_vals = [dlg.result[c] or None for c in self.columns[1:]] + [pk_val]
        try:
            self.cursor.execute(sql, new_vals)
            self.conn.commit()
            self.refresh()
            self.set_status("Record updated.", ok=True)
        except Error as e:
            messagebox.showerror("Update Error", str(e))

    def delete_record(self):
        values = self._selected_values()
        if values is None:
            return
        pk_col = self.columns[0]
        pk_val = values[0]
        if not messagebox.askyesno(
            "Confirm Delete",
            f"Delete record where {pk_col} = {pk_val}?\n\nThis cannot be undone."
        ):
            return
        sql = f"DELETE FROM `{self.current_table.get()}` WHERE `{pk_col}` = %s"
        try:
            self.cursor.execute(sql, (pk_val,))
            self.conn.commit()
            self.refresh()
            self.set_status("Record deleted.", ok=True)
        except Error as e:
            messagebox.showerror("Delete Error", str(e))

    # ── Export ────────────────────────────────────────────────────────────────

    def export_data(self):
        if not self.all_rows:
            messagebox.showwarning("No Data", "No data to export.")
            return

        if EXCEL_OK:
            filetypes = [("Excel file", "*.xlsx"), ("CSV file", "*.csv")]
        else:
            filetypes = [("CSV file", "*.csv")]

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx" if EXCEL_OK else ".csv",
            filetypes=filetypes,
            initialfile=self.current_table.get(),
            title="Export Table"
        )
        if not path:
            return

        try:
            if path.endswith(".xlsx") and EXCEL_OK:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = self.current_table.get()
                ws.append(self.columns)
                for row in self.all_rows:
                    ws.append([str(v) if v is not None else "" for v in row])
                from openpyxl.styles import Font, PatternFill
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="0D1120")
                wb.save(path)
            else:
                with open(path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(self.columns)
                    writer.writerows(self.all_rows)

            self.set_status(f"Exported: {os.path.basename(path)}", ok=True)
            messagebox.showinfo("Export Done", f"File saved:\n{path}")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

    # ── Dashboard ─────────────────────────────────────────────────────────────

    def open_dashboard(self):
        self.ensure_connection()
        if not self.conn:
            return
        DashboardWindow(self, self.cursor)

    # ── Status bar ────────────────────────────────────────────────────────────

    def set_status(self, msg, ok=True):
        self.status_text.set(msg)
        self.lbl_status.config(fg=ACCENT if ok else BTN_DEL)


# ─── Entry Point ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app = SkillTrackApp()
    app.mainloop()
